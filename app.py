from __future__ import annotations

import csv
import base64
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import sqlite3
from datetime import datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import psycopg
    from psycopg.rows import dict_row
except ImportError:  # Local SQLite mode does not need psycopg.
    psycopg = None
    dict_row = None


ROOT = Path(__file__).resolve().parent
CSV_PATH = ROOT / "DATOS_IMPORTADOS" / "inventario_unificado.csv"
DB_PATH = ROOT / "inventario.db"
WEB_DIR = ROOT / "web"
DATABASE_URL = os.environ.get("DATABASE_URL", "")
IS_POSTGRES = DATABASE_URL.startswith(("postgres://", "postgresql://"))


def sql_params(sql: str) -> str:
    return sql.replace("?", "%s") if IS_POSTGRES else sql


def connect():
    if IS_POSTGRES:
        if psycopg is None:
            raise RuntimeError("Falta instalar psycopg para usar PostgreSQL")
        return psycopg.connect(DATABASE_URL, row_factory=dict_row)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def execute(conn, sql: str, params: tuple | list = ()):
    if IS_POSTGRES:
        cur = conn.cursor()
        cur.execute(sql_params(sql), params)
        return cur
    return conn.execute(sql, params)


def executemany(conn, sql: str, rows: list[tuple]) -> None:
    if IS_POSTGRES:
        cur = conn.cursor()
        cur.executemany(sql_params(sql), rows)
        return
    conn.executemany(sql, rows)


def run_script(conn, script: str) -> None:
    if not IS_POSTGRES:
        conn.executescript(script)
        return
    for statement in script.split(";"):
        statement = statement.strip()
        if statement:
            execute(conn, statement)


def begin_write(conn) -> None:
    execute(conn, "BEGIN" if IS_POSTGRES else "BEGIN IMMEDIATE")


def to_int(value: object) -> int:
    text = "" if value is None else str(value).strip().replace(",", ".")
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def now_text() -> str:
    return datetime.now().isoformat(timespec="seconds")


def row_to_dict(row) -> dict:
    return dict(row) if row is not None else {}


def normalized_stock(row) -> dict:
    product = row_to_dict(row)
    unidad = to_int(product.get("stock_unidad"))
    deposito = to_int(product.get("stock_deposito"))
    if unidad == 0 and deposito == 0 and to_int(product.get("stock")) != 0:
        unidad = to_int(product.get("stock"))
    product["stock_unidad"] = unidad
    product["stock_deposito"] = deposito
    product["stock"] = unidad + deposito
    product["stock_total"] = unidad + deposito
    return product


def compact_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def pick(row: dict, *names: str) -> object:
    normalized = {compact_key(key): value for key, value in row.items()}
    for name in names:
        key = compact_key(name)
        if key in normalized:
            return normalized[key]
    return ""


def hash_password(password: str, salt: str | None = None) -> str:
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 120_000)
    return f"{salt}${digest.hex()}"


def verify_password(password: str, stored: str) -> bool:
    try:
        salt, expected = stored.split("$", 1)
    except ValueError:
        return False
    candidate = hash_password(password, salt).split("$", 1)[1]
    return hmac.compare_digest(candidate, expected)


def table_columns(conn, table: str) -> set[str]:
    if IS_POSTGRES:
        rows = execute(
            conn,
            """
            SELECT column_name AS name
            FROM information_schema.columns
            WHERE table_name = ?
            """,
            (table,),
        ).fetchall()
        return {row["name"] for row in rows}
    return {row["name"] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}


def migrate_users_table(conn: sqlite3.Connection) -> None:
    if IS_POSTGRES:
        return
    schema = execute(conn, 
        "SELECT sql FROM sqlite_master WHERE type = 'table' AND name = 'usuarios'"
    ).fetchone()
    if schema is None or "'dueno'" not in (schema["sql"] or ""):
        return

    execute(conn, "ALTER TABLE usuarios RENAME TO usuarios_old")
    execute(conn, 
        """
        CREATE TABLE usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            usuario TEXT NOT NULL UNIQUE,
            nombre TEXT NOT NULL,
            password_hash TEXT NOT NULL,
            rol TEXT NOT NULL CHECK (rol IN ('administrador', 'empleado')),
            activo INTEGER NOT NULL DEFAULT 1
        );
        """
    )
    execute(conn, 
        """
        INSERT INTO usuarios (id, usuario, nombre, password_hash, rol, activo)
        SELECT id,
               CASE WHEN usuario = 'dueno' THEN 'admin' ELSE usuario END,
               CASE WHEN rol = 'dueno' THEN 'Administrador' ELSE nombre END,
               password_hash,
               CASE WHEN rol = 'dueno' THEN 'administrador' ELSE rol END,
               activo
        FROM usuarios_old
        """
    )
    execute(conn, "DROP TABLE usuarios_old")


def init_db() -> None:
    conn = connect()
    id_type = "SERIAL PRIMARY KEY" if IS_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
    script = f"""
        CREATE TABLE IF NOT EXISTS usuarios (
            id {id_type},
            usuario TEXT NOT NULL UNIQUE,
            nombre TEXT NOT NULL,
            password_hash TEXT NOT NULL,
            rol TEXT NOT NULL CHECK (rol IN ('administrador', 'empleado')),
            activo INTEGER NOT NULL DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS sesiones (
            token TEXT PRIMARY KEY,
            usuario_id INTEGER NOT NULL,
            creado TEXT NOT NULL,
            FOREIGN KEY (usuario_id) REFERENCES usuarios(id)
        );

        CREATE TABLE IF NOT EXISTS productos (
            id {id_type},
            archivo_origen TEXT NOT NULL,
            hoja_origen TEXT NOT NULL,
            fila_origen INTEGER NOT NULL,
            codigo_item TEXT,
            codigo_barras TEXT,
            codigo_articulo TEXT,
            producto TEXT,
            marca TEXT,
            proveedor TEXT,
            departamento TEXT,
            descripcion TEXT,
            aplicacion TEXT,
            stock INTEGER NOT NULL DEFAULT 0,
            stock_unidad INTEGER NOT NULL DEFAULT 0,
            stock_deposito INTEGER NOT NULL DEFAULT 0,
            ubicacion TEXT,
            stock_minimo INTEGER NOT NULL DEFAULT 0,
            fecha_actualizacion TEXT,
            UNIQUE (archivo_origen, hoja_origen, fila_origen)
        );

        CREATE TABLE IF NOT EXISTS movimientos (
            id {id_type},
            producto_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            tipo TEXT NOT NULL,
            cantidad INTEGER NOT NULL,
            stock_anterior INTEGER NOT NULL,
            stock_nuevo INTEGER NOT NULL,
            nota TEXT,
            motivo TEXT,
            usuario_id INTEGER,
            usuario_nombre TEXT,
            FOREIGN KEY (producto_id) REFERENCES productos(id)
        );

        CREATE TABLE IF NOT EXISTS ventas (
            id {id_type},
            fecha TEXT NOT NULL,
            usuario_id INTEGER,
            usuario_nombre TEXT,
            cliente TEXT,
            nota TEXT,
            total_items INTEGER NOT NULL DEFAULT 0,
            total_unidades INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS venta_items (
            id {id_type},
            venta_id INTEGER NOT NULL,
            producto_id INTEGER NOT NULL,
            codigo_item TEXT,
            producto TEXT,
            marca TEXT,
            descripcion TEXT,
            cantidad INTEGER NOT NULL,
            stock_anterior INTEGER NOT NULL,
            stock_nuevo INTEGER NOT NULL,
            FOREIGN KEY (venta_id) REFERENCES ventas(id),
            FOREIGN KEY (producto_id) REFERENCES productos(id)
        );
        """
    run_script(conn, script)
    migrate_users_table(conn)

    movement_columns = table_columns(conn, "movimientos")
    if "usuario_id" not in movement_columns:
        execute(conn, "ALTER TABLE movimientos ADD COLUMN usuario_id INTEGER")
    if "usuario_nombre" not in movement_columns:
        execute(conn, "ALTER TABLE movimientos ADD COLUMN usuario_nombre TEXT")

    product_columns = table_columns(conn, "productos")
    if "ubicacion" not in product_columns:
        execute(conn, "ALTER TABLE productos ADD COLUMN ubicacion TEXT")
    if "stock_minimo" not in product_columns:
        execute(conn, "ALTER TABLE productos ADD COLUMN stock_minimo INTEGER NOT NULL DEFAULT 0")
    if "stock_unidad" not in product_columns:
        execute(conn, "ALTER TABLE productos ADD COLUMN stock_unidad INTEGER NOT NULL DEFAULT 0")
    if "stock_deposito" not in product_columns:
        execute(conn, "ALTER TABLE productos ADD COLUMN stock_deposito INTEGER NOT NULL DEFAULT 0")
    if "aplicacion" not in product_columns:
        execute(conn, "ALTER TABLE productos ADD COLUMN aplicacion TEXT")
    if "fecha_actualizacion" not in product_columns:
        execute(conn, "ALTER TABLE productos ADD COLUMN fecha_actualizacion TEXT")

    movement_columns = table_columns(conn, "movimientos")
    if "motivo" not in movement_columns:
        execute(conn, "ALTER TABLE movimientos ADD COLUMN motivo TEXT")

    execute(
        conn,
        """
        UPDATE productos
        SET stock_unidad = stock
        WHERE COALESCE(stock_unidad, 0) = 0
          AND COALESCE(stock_deposito, 0) = 0
          AND COALESCE(stock, 0) <> 0
        """,
    )
    execute(
        conn,
        """
        UPDATE productos
        SET stock = COALESCE(stock_unidad, 0) + COALESCE(stock_deposito, 0)
        """,
    )

    execute(conn, "UPDATE usuarios SET usuario = 'admin' WHERE usuario = 'dueno'")
    execute(conn, "UPDATE usuarios SET rol = 'administrador' WHERE rol = 'dueno'")
    execute(conn, "UPDATE usuarios SET nombre = 'Administrador' WHERE usuario = 'admin'")

    user_count_row = execute(conn, "SELECT COUNT(*) AS count FROM usuarios").fetchone()
    user_count = user_count_row["count"] if IS_POSTGRES else user_count_row[0]
    if user_count == 0:
        executemany(
            conn,
            """
            INSERT INTO usuarios (usuario, nombre, password_hash, rol)
            VALUES (?, ?, ?, ?)
            """,
            [
                ("admin", "Administrador", hash_password("admin123"), "administrador"),
                ("empleado", "Empleado", hash_password("empleado123"), "empleado"),
            ],
        )

    count_row = execute(conn, "SELECT COUNT(*) AS count FROM productos").fetchone()
    count = count_row["count"] if IS_POSTGRES else count_row[0]
    if count == 0:
        import_csv(conn)

    conn.commit()
    conn.close()


def import_csv(conn: sqlite3.Connection) -> None:
    if not CSV_PATH.exists():
        raise FileNotFoundError(f"No existe {CSV_PATH}")

    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        rows = []
        for row in reader:
            stock_unidad = to_int(row.get("unidad") or row.get("stock_unidad") or row.get("stock_calculado") or row.get("stock"))
            stock_deposito = to_int(row.get("deposito") or row.get("stock_deposito") or row.get("depósito"))
            stock = stock_unidad + stock_deposito
            rows.append(
                (
                    row.get("archivo_origen", ""),
                    row.get("hoja_origen", ""),
                    to_int(row.get("fila_origen")),
                    row.get("codigo_item", ""),
                    row.get("codigo_barras", ""),
                    row.get("codigo_articulo", ""),
                    row.get("producto", ""),
                    row.get("marca", ""),
                    row.get("proveedor", ""),
                    row.get("departamento", ""),
                    row.get("descripcion", ""),
                    stock,
                    stock_unidad,
                    stock_deposito,
                    now_text(),
                )
            )

    insert_sql = """
        INSERT INTO productos (
            archivo_origen, hoja_origen, fila_origen, codigo_item, codigo_barras,
            codigo_articulo, producto, marca, proveedor, departamento, descripcion,
            stock, stock_unidad, stock_deposito, fecha_actualizacion
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """
    if IS_POSTGRES:
        insert_sql += " ON CONFLICT (archivo_origen, hoja_origen, fila_origen) DO NOTHING"
    else:
        insert_sql = insert_sql.replace("INSERT INTO", "INSERT OR IGNORE INTO", 1)
    executemany(conn, insert_sql, rows)


def json_response(handler: SimpleHTTPRequestHandler, data: object, status: int = 200) -> None:
    payload = json.dumps(data, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(payload)))
    handler.end_headers()
    handler.wfile.write(payload)


def parse_cookies(cookie_header: str | None) -> dict[str, str]:
    cookies = {}
    for part in (cookie_header or "").split(";"):
        if "=" not in part:
            continue
        key, value = part.strip().split("=", 1)
        cookies[key] = value
    return cookies


class AppHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(WEB_DIR), **kwargs)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/me":
            self.handle_me()
            return
        if parsed.path == "/api/productos":
            if not self.require_user():
                return
            self.handle_products(parsed.query)
            return
        if parsed.path == "/api/productos/buscar":
            if not self.require_user():
                return
            self.handle_products(parsed.query)
            return
        if parsed.path == "/api/filtros":
            if not self.require_user():
                return
            self.handle_filters()
            return
        if parsed.path == "/api/resumen":
            if not self.require_user():
                return
            self.handle_summary()
            return
        if parsed.path == "/api/inventario/resumen":
            if not self.require_user():
                return
            self.handle_inventory_summary(parsed.query)
            return
        if parsed.path == "/api/dashboard":
            if not self.require_user():
                return
            self.handle_dashboard()
            return
        if parsed.path == "/api/movimientos":
            if not self.require_user(role="administrador"):
                return
            self.handle_movements(parsed.query)
            return
        if parsed.path == "/api/ventas/resumen":
            if not self.require_user():
                return
            self.handle_sales_summary()
            return
        if parsed.path == "/":
            self.path = "/index.html"
        super().do_GET()

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/login":
            self.handle_login()
            return
        if parsed.path == "/api/logout":
            self.handle_logout()
            return
        if parsed.path == "/api/venta":
            if not self.require_user():
                return
            self.handle_sale()
            return
        if parsed.path == "/api/ventas":
            if not self.require_user():
                return
            self.handle_cart_sale()
            return
        if parsed.path == "/api/venta-carrito":
            if not self.require_user():
                return
            self.handle_cart_sale()
            return
        if parsed.path == "/api/producto/actualizar":
            if not self.require_user(role="administrador"):
                return
            self.handle_product_update()
            return
        if parsed.path == "/api/stock/ajustar":
            if not self.require_user():
                return
            self.handle_stock_adjust()
            return
        if parsed.path == "/api/stock/importar-excel":
            if not self.require_user(role="administrador"):
                return
            self.handle_excel_import()
            return
        json_response(self, {"error": "Ruta no encontrada"}, 404)

    def read_json(self) -> dict:
        length = int(self.headers.get("Content-Length", "0"))
        payload = self.rfile.read(length)
        try:
            raw = payload.decode("utf-8")
        except UnicodeDecodeError:
            raw = payload.decode("latin-1")
        return json.loads(raw or "{}")

    def current_user(self) -> sqlite3.Row | None:
        token = parse_cookies(self.headers.get("Cookie")).get("session")
        if not token:
            return None
        conn = connect()
        user = execute(conn, 
            """
            SELECT u.id, u.usuario, u.nombre, u.rol
            FROM sesiones s
            JOIN usuarios u ON u.id = s.usuario_id
            WHERE s.token = ? AND u.activo = 1
            """,
            (token,),
        ).fetchone()
        conn.close()
        return user

    def require_user(self, role: str | None = None) -> sqlite3.Row | None:
        user = self.current_user()
        if user is None:
            json_response(self, {"error": "Sesion requerida"}, 401)
            return None
        if role and user["rol"] != role:
            json_response(self, {"error": "No tenes permiso para esta accion"}, 403)
            return None
        return user

    def handle_me(self) -> None:
        user = self.current_user()
        if user is None:
            json_response(self, {"user": None}, 401)
            return
        json_response(self, {"user": dict(user)})

    def handle_login(self) -> None:
        data = self.read_json()
        username = str(data.get("usuario", "")).strip().lower()
        password = str(data.get("password", ""))
        conn = connect()
        user = execute(conn, 
            "SELECT id, usuario, nombre, password_hash, rol FROM usuarios WHERE usuario = ? AND activo = 1",
            (username,),
        ).fetchone()
        if user is None or not verify_password(password, user["password_hash"]):
            conn.close()
            json_response(self, {"error": "Usuario o contrasena incorrectos"}, 401)
            return

        token = secrets.token_urlsafe(32)
        execute(conn, 
            "INSERT INTO sesiones (token, usuario_id, creado) VALUES (?, ?, ?)",
            (token, user["id"], datetime.now().isoformat(timespec="seconds")),
        )
        conn.commit()
        conn.close()

        payload = json.dumps(
            {"user": {"id": user["id"], "usuario": user["usuario"], "nombre": user["nombre"], "rol": user["rol"]}},
            ensure_ascii=False,
        ).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Set-Cookie", f"session={token}; Path=/; HttpOnly; SameSite=Lax")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def handle_logout(self) -> None:
        token = parse_cookies(self.headers.get("Cookie")).get("session")
        if token:
            conn = connect()
            execute(conn, "DELETE FROM sesiones WHERE token = ?", (token,))
            conn.commit()
            conn.close()
        payload = b'{"ok": true}'
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Set-Cookie", "session=; Path=/; Max-Age=0; HttpOnly; SameSite=Lax")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def handle_products(self, query: str) -> None:
        params = parse_qs(query)
        search = params.get("q", [""])[0].strip()
        proveedor = params.get("proveedor", [""])[0].strip()
        categoria = params.get("categoria", [""])[0].strip()
        limit = min(to_int(params.get("limit", ["100"])[0]) or 100, 300)

        sql = """
            SELECT id, codigo_item, codigo_barras, codigo_articulo, producto, marca,
                   proveedor, departamento, descripcion, aplicacion,
                   COALESCE(stock_unidad, stock, 0) AS stock_unidad,
                   COALESCE(stock_deposito, 0) AS stock_deposito,
                   COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0) AS stock,
                   ubicacion, stock_minimo, hoja_origen, fecha_actualizacion
            FROM productos
        """
        values: list[object] = []
        if search:
            terms = [term for term in search.split() if term]
            clauses = []
            for term in terms:
                like = f"%{term}%"
                clauses.append(
                    """
                    (LOWER(COALESCE(codigo_item, '')) LIKE ? OR LOWER(COALESCE(codigo_barras, '')) LIKE ?
                     OR LOWER(COALESCE(codigo_articulo, '')) LIKE ? OR LOWER(COALESCE(producto, '')) LIKE ?
                     OR LOWER(COALESCE(marca, '')) LIKE ? OR LOWER(COALESCE(proveedor, '')) LIKE ?
                     OR LOWER(COALESCE(departamento, '')) LIKE ? OR LOWER(COALESCE(descripcion, '')) LIKE ?
                     OR LOWER(COALESCE(aplicacion, '')) LIKE ?)
                    """
                )
                values.extend([like.lower()] * 9)
            sql += " WHERE " + " AND ".join(clauses)
        filters = []
        if proveedor:
            filters.append("proveedor = ?")
            values.append(proveedor)
        if categoria:
            filters.append("departamento = ?")
            values.append(categoria)
        if filters:
            sql += (" AND " if " WHERE " in sql else " WHERE ") + " AND ".join(filters)
        sql += " ORDER BY stock DESC, codigo_item LIMIT ?"
        values.append(limit)

        conn = connect()
        products = [normalized_stock(row) for row in execute(conn, sql, values).fetchall()]
        conn.close()
        json_response(self, {"productos": products})

    def handle_product_update(self) -> None:
        data = self.read_json()
        product_id = to_int(data.get("id"))
        if product_id <= 0:
            json_response(self, {"error": "Producto invalido"}, 400)
            return

        fields = {
            "codigo_item": str(data.get("codigo_item", "")).strip(),
            "codigo_barras": str(data.get("codigo_barras", "")).strip(),
            "codigo_articulo": str(data.get("codigo_articulo", "")).strip(),
            "producto": str(data.get("producto", "")).strip(),
            "marca": str(data.get("marca", "")).strip(),
            "proveedor": str(data.get("proveedor", "")).strip(),
            "departamento": str(data.get("departamento", "")).strip(),
            "descripcion": str(data.get("descripcion", "")).strip(),
            "aplicacion": str(data.get("aplicacion", "")).strip(),
            "ubicacion": str(data.get("ubicacion", "")).strip(),
            "stock_minimo": to_int(data.get("stock_minimo")),
        }

        conn = connect()
        existing = execute(conn, "SELECT id FROM productos WHERE id = ?", (product_id,)).fetchone()
        if existing is None:
            conn.close()
            json_response(self, {"error": "Producto no encontrado"}, 404)
            return

        execute(conn, 
            """
            UPDATE productos
            SET codigo_item = ?, codigo_barras = ?, codigo_articulo = ?, producto = ?,
                marca = ?, proveedor = ?, departamento = ?, descripcion = ?,
                aplicacion = ?, ubicacion = ?, stock_minimo = ?, fecha_actualizacion = ?
            WHERE id = ?
            """,
            (
                fields["codigo_item"],
                fields["codigo_barras"],
                fields["codigo_articulo"],
                fields["producto"],
                fields["marca"],
                fields["proveedor"],
                fields["departamento"],
                fields["descripcion"],
                fields["aplicacion"],
                fields["ubicacion"],
                fields["stock_minimo"],
                now_text(),
                product_id,
            ),
        )
        conn.commit()
        updated = execute(conn, 
            """
            SELECT id, codigo_item, codigo_barras, codigo_articulo, producto, marca,
                   proveedor, departamento, descripcion, aplicacion,
                   COALESCE(stock_unidad, stock, 0) AS stock_unidad,
                   COALESCE(stock_deposito, 0) AS stock_deposito,
                   COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0) AS stock,
                   ubicacion, stock_minimo, hoja_origen, fecha_actualizacion
            FROM productos
            WHERE id = ?
            """,
            (product_id,),
        ).fetchone()
        conn.close()
        json_response(self, {"ok": True, "producto": normalized_stock(updated)})

    def handle_filters(self) -> None:
        conn = connect()
        proveedores = [
            row["proveedor"]
            for row in execute(conn, 
                """
                SELECT proveedor, COUNT(*) AS total
                FROM productos
                WHERE proveedor IS NOT NULL AND TRIM(proveedor) <> ''
                GROUP BY proveedor
                ORDER BY total DESC, proveedor
                LIMIT 200
                """
            ).fetchall()
        ]
        categorias = [
            row["departamento"]
            for row in execute(conn, 
                """
                SELECT departamento, COUNT(*) AS total
                FROM productos
                WHERE departamento IS NOT NULL AND TRIM(departamento) <> ''
                GROUP BY departamento
                ORDER BY total DESC, departamento
                LIMIT 200
                """
            ).fetchall()
        ]
        conn.close()
        json_response(self, {"proveedores": proveedores, "categorias": categorias})

    def handle_summary(self) -> None:
        conn = connect()
        row = execute(conn, 
            """
            SELECT
                COUNT(*) AS productos,
                SUM(CASE WHEN (COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)) > 0 THEN 1 ELSE 0 END) AS con_stock,
                SUM(COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)) AS unidades,
                COUNT(DISTINCT departamento) AS departamentos
            FROM productos
            """
        ).fetchone()
        conn.close()
        json_response(self, dict(row))

    def handle_dashboard(self) -> None:
        conn = connect()
        kpis = execute(conn, 
            """
            SELECT
                COUNT(*) AS productos_total,
                SUM(CASE WHEN (COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)) > 0 THEN 1 ELSE 0 END) AS productos_con_stock,
                SUM(CASE WHEN (COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)) <= 0 THEN 1 ELSE 0 END) AS productos_sin_stock,
                SUM(CASE WHEN stock_minimo > 0 AND (COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)) <= stock_minimo THEN 1 ELSE 0 END) AS stock_critico,
                COALESCE(SUM(COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)), 0) AS unidades_disponibles,
                COUNT(DISTINCT proveedor) AS proveedores,
                COUNT(DISTINCT departamento) AS categorias
            FROM productos
            """
        ).fetchone()
        sales = execute(conn, 
            """
            SELECT
                COUNT(*) AS ventas_total,
                COALESCE(SUM(cantidad), 0) AS unidades_vendidas
            FROM movimientos
            WHERE tipo = 'venta'
            """
        ).fetchone()
        top_departments = [
            dict(row)
            for row in execute(conn, 
                """
                SELECT COALESCE(NULLIF(TRIM(departamento), ''), 'Sin categoria') AS nombre,
                       COUNT(*) AS productos,
                       COALESCE(SUM(COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)), 0) AS unidades
                FROM productos
                GROUP BY nombre
                ORDER BY productos DESC
                LIMIT 8
                """
            ).fetchall()
        ]
        critical_products = [
            dict(row)
            for row in execute(conn, 
                """
                SELECT id, codigo_item, producto, marca, proveedor, departamento,
                       descripcion, COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0) AS stock, stock_minimo, ubicacion
                FROM productos
                WHERE (COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)) <= 0
                   OR (stock_minimo > 0 AND (COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0)) <= stock_minimo)
                ORDER BY stock ASC, codigo_item
                LIMIT 10
                """
            ).fetchall()
        ]
        top_sales = [
            dict(row)
            for row in execute(conn, 
                """
                SELECT p.id, p.codigo_item, p.producto, p.marca, p.descripcion,
                       SUM(m.cantidad) AS unidades
                FROM movimientos m
                JOIN productos p ON p.id = m.producto_id
                WHERE m.tipo = 'venta'
                GROUP BY p.id, p.codigo_item, p.producto, p.marca, p.descripcion
                ORDER BY unidades DESC
                LIMIT 8
                """
            ).fetchall()
        ]
        conn.close()
        json_response(
            self,
            {
                "kpis": {**dict(kpis), **dict(sales)},
                "categorias_principales": top_departments,
                "productos_criticos": critical_products,
                "productos_mayor_salida": top_sales,
            },
        )

    def handle_movements(self, query: str) -> None:
        params = parse_qs(query)
        product_id = to_int(params.get("producto_id", ["0"])[0])
        conn = connect()
        rows = execute(conn, 
            """
            SELECT m.id, m.fecha, m.tipo, m.cantidad, m.stock_anterior, m.stock_nuevo,
                   COALESCE(m.motivo, m.nota, '') AS motivo, m.nota, m.usuario_nombre,
                   p.codigo_item, p.producto, p.marca, p.descripcion
            FROM movimientos m
            JOIN productos p ON p.id = m.producto_id
            WHERE (? = 0 OR m.producto_id = ?)
            ORDER BY m.id DESC
            LIMIT 100
            """,
            (product_id, product_id),
        ).fetchall()
        conn.close()
        json_response(self, {"movimientos": [dict(row) for row in rows]})

    def handle_sales_summary(self) -> None:
        conn = connect()
        totals = execute(conn, 
            """
            SELECT
                COUNT(*) AS operaciones,
                COALESCE(SUM(cantidad), 0) AS unidades
            FROM movimientos
            WHERE tipo = 'venta'
            """
        ).fetchone()
        top_products = [
            dict(row)
            for row in execute(conn, 
                """
                SELECT p.id, p.codigo_item, p.producto, p.marca, p.descripcion,
                       COALESCE(SUM(m.cantidad), 0) AS unidades,
                       COUNT(*) AS operaciones
                FROM movimientos m
                JOIN productos p ON p.id = m.producto_id
                WHERE m.tipo = 'venta'
                GROUP BY p.id, p.codigo_item, p.producto, p.marca, p.descripcion
                ORDER BY unidades DESC, operaciones DESC, p.codigo_item
                LIMIT 10
                """
            ).fetchall()
        ]
        recent_sales = [
            dict(row)
            for row in execute(conn, 
                """
                SELECT m.id, m.fecha, m.cantidad, m.stock_anterior, m.stock_nuevo,
                       m.usuario_nombre, m.nota, p.codigo_item, p.producto, p.marca, p.descripcion
                FROM movimientos m
                JOIN productos p ON p.id = m.producto_id
                WHERE m.tipo = 'venta'
                ORDER BY m.id DESC
                LIMIT 10
                """
            ).fetchall()
        ]
        conn.close()
        json_response(
            self,
            {
                "operaciones": totals["operaciones"],
                "unidades": totals["unidades"],
                "productos_mayor_salida": top_products,
                "ultimas_ventas": recent_sales,
            },
        )

    def apply_stock_out(self, conn, product_id: int, quantity: int) -> tuple[dict, int, int, int, int]:
        product = execute(
            conn,
            """
            SELECT id, codigo_item, producto, marca, descripcion,
                   COALESCE(stock_unidad, stock, 0) AS stock_unidad,
                   COALESCE(stock_deposito, 0) AS stock_deposito
            FROM productos
            WHERE id = ?
            """,
            (product_id,),
        ).fetchone()
        if product is None:
            raise ValueError("Producto no encontrado")

        product_dict = normalized_stock(product)
        previous_total = product_dict["stock_total"]
        if previous_total < quantity:
            raise ValueError(f"No hay stock suficiente para {product_dict.get('codigo_item') or 'el producto'}")

        previous_unit = product_dict["stock_unidad"]
        previous_deposit = product_dict["stock_deposito"]
        take_unit = min(previous_unit, quantity)
        take_deposit = quantity - take_unit
        new_unit = previous_unit - take_unit
        new_deposit = previous_deposit - take_deposit
        new_total = new_unit + new_deposit

        execute(
            conn,
            """
            UPDATE productos
            SET stock_unidad = ?, stock_deposito = ?, stock = ?, fecha_actualizacion = ?
            WHERE id = ?
            """,
            (new_unit, new_deposit, new_total, now_text(), product_id),
        )
        return product_dict, previous_total, new_total, new_unit, new_deposit

    def register_movement(
        self,
        conn,
        product_id: int,
        kind: str,
        quantity: int,
        previous_stock: int,
        new_stock: int,
        reason: str,
        user: dict,
    ) -> None:
        execute(
            conn,
            """
            INSERT INTO movimientos (
                producto_id, fecha, tipo, cantidad, stock_anterior, stock_nuevo,
                nota, motivo, usuario_id, usuario_nombre
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                product_id,
                now_text(),
                kind,
                quantity,
                previous_stock,
                new_stock,
                reason,
                reason,
                user["id"],
                user["nombre"],
            ),
        )

    def handle_sale(self) -> None:
        user = self.require_user()
        if user is None:
            return
        data = self.read_json()
        product_id = to_int(data.get("producto_id"))
        quantity = to_int(data.get("cantidad"))
        note = str(data.get("nota", "")).strip()

        if product_id <= 0 or quantity <= 0:
            json_response(self, {"error": "Producto y cantidad son obligatorios"}, 400)
            return

        conn = connect()
        try:
            begin_write(conn)
            _, previous, new_stock, new_unit, new_deposit = self.apply_stock_out(conn, product_id, quantity)
            self.register_movement(conn, product_id, "venta", quantity, previous, new_stock, note or "venta", user)
            conn.commit()
        except ValueError as exc:
            conn.rollback()
            json_response(self, {"error": str(exc)}, 400)
            return
        finally:
            conn.close()

        json_response(
            self,
            {"ok": True, "stock_nuevo": new_stock, "stock_unidad": new_unit, "stock_deposito": new_deposit},
        )

    def handle_cart_sale(self) -> None:
        user = self.require_user()
        if user is None:
            return
        data = self.read_json()
        items = data.get("items", [])
        cliente = str(data.get("cliente", "")).strip()
        note = str(data.get("nota", "")).strip()

        if not isinstance(items, list) or not items:
            json_response(self, {"error": "La venta no tiene productos"}, 400)
            return

        quantities: dict[int, int] = {}
        for item in items:
            product_id = to_int(item.get("producto_id"))
            quantity = to_int(item.get("cantidad"))
            if product_id <= 0 or quantity <= 0:
                json_response(self, {"error": "Hay productos o cantidades invalidas"}, 400)
                return
            quantities[product_id] = quantities.get(product_id, 0) + quantity

        conn = connect()
        try:
            begin_write(conn)
            placeholders = ",".join("?" for _ in quantities)
            products = {
                row["id"]: normalized_stock(row)
                for row in execute(conn, 
                    f"""
                    SELECT id, codigo_item, producto, marca, descripcion,
                           COALESCE(stock_unidad, stock, 0) AS stock_unidad,
                           COALESCE(stock_deposito, 0) AS stock_deposito
                    FROM productos
                    WHERE id IN ({placeholders})
                    """,
                    tuple(quantities.keys()),
                ).fetchall()
            }

            if len(products) != len(quantities):
                raise ValueError("Uno o mas productos no existen")

            for product_id, quantity in quantities.items():
                if int(products[product_id]["stock_total"]) < quantity:
                    raise ValueError(f"No hay stock suficiente para {products[product_id]['codigo_item']}")

            now = datetime.now().isoformat(timespec="seconds")
            total_items = len(quantities)
            total_units = sum(quantities.values())
            insert_sale_sql = """
                INSERT INTO ventas (fecha, usuario_id, usuario_nombre, cliente, nota, total_items, total_unidades)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """
            if IS_POSTGRES:
                insert_sale_sql += " RETURNING id"
            cursor = execute(
                conn,
                insert_sale_sql,
                (now, user["id"], user["nombre"], cliente, note, total_items, total_units),
            )
            sale_id = cursor.fetchone()["id"] if IS_POSTGRES else cursor.lastrowid

            response_items = []
            for product_id, quantity in quantities.items():
                product, previous, new_stock, new_unit, new_deposit = self.apply_stock_out(conn, product_id, quantity)
                execute(conn, 
                    """
                    INSERT INTO venta_items (
                        venta_id, producto_id, codigo_item, producto, marca, descripcion,
                        cantidad, stock_anterior, stock_nuevo
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        sale_id,
                        product_id,
                        product["codigo_item"],
                        product["producto"],
                        product["marca"],
                        product["descripcion"],
                        quantity,
                        previous,
                        new_stock,
                    ),
                )
                execute(conn, 
                    """
                    INSERT INTO movimientos (
                        producto_id, fecha, tipo, cantidad, stock_anterior, stock_nuevo,
                        nota, motivo, usuario_id, usuario_nombre
                    )
                    VALUES (?, ?, 'venta', ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        product_id,
                        now,
                        quantity,
                        previous,
                        new_stock,
                        f"Venta #{sale_id}" + (f" - {note}" if note else ""),
                        f"Venta #{sale_id}" + (f" - {note}" if note else ""),
                        user["id"],
                        user["nombre"],
                    ),
                )
                response_items.append(
                    {
                        "producto_id": product_id,
                        "codigo_item": product["codigo_item"],
                        "cantidad": quantity,
                        "stock_nuevo": new_stock,
                        "stock_unidad": new_unit,
                        "stock_deposito": new_deposit,
                    }
                )

            conn.commit()
        except ValueError as exc:
            conn.rollback()
            json_response(self, {"error": str(exc)}, 400)
            return
        finally:
            conn.close()

        json_response(
            self,
            {
                "ok": True,
                "venta_id": sale_id,
                "total_items": total_items,
                "total_unidades": total_units,
                "items": response_items,
            },
        )

    def handle_stock_adjust(self) -> None:
        user = self.require_user()
        if user is None:
            return
        data = self.read_json()
        product_id = to_int(data.get("producto_id") or data.get("id"))
        stock_unidad = to_int(data.get("stock_unidad"))
        stock_deposito = to_int(data.get("stock_deposito"))
        reason = str(data.get("motivo", "")).strip() or "ajuste manual"
        if product_id <= 0:
            json_response(self, {"error": "Seleccioná un producto válido"}, 400)
            return
        if stock_unidad < 0 or stock_deposito < 0:
            json_response(self, {"error": "El stock no puede ser negativo"}, 400)
            return

        conn = connect()
        try:
            begin_write(conn)
            product = execute(
                conn,
                """
                SELECT id, COALESCE(stock_unidad, stock, 0) AS stock_unidad,
                       COALESCE(stock_deposito, 0) AS stock_deposito
                FROM productos
                WHERE id = ?
                """,
                (product_id,),
            ).fetchone()
            if product is None:
                raise ValueError("Producto no encontrado")
            previous = normalized_stock(product)["stock_total"]
            new_stock = stock_unidad + stock_deposito
            execute(
                conn,
                """
                UPDATE productos
                SET stock_unidad = ?, stock_deposito = ?, stock = ?, fecha_actualizacion = ?
                WHERE id = ?
                """,
                (stock_unidad, stock_deposito, new_stock, now_text(), product_id),
            )
            self.register_movement(
                conn,
                product_id,
                "ajuste_stock",
                new_stock - previous,
                previous,
                new_stock,
                reason,
                user,
            )
            conn.commit()
        except ValueError as exc:
            conn.rollback()
            json_response(self, {"error": str(exc)}, 400)
            return
        finally:
            conn.close()
        json_response(
            self,
            {
                "ok": True,
                "stock_unidad": stock_unidad,
                "stock_deposito": stock_deposito,
                "stock_nuevo": new_stock,
            },
        )

    def handle_inventory_summary(self, query: str) -> None:
        params = parse_qs(query)
        departamento = params.get("departamento", [""])[0].strip()
        marca = params.get("marca", [""])[0].strip()
        proveedor = params.get("proveedor", [""])[0].strip()
        estado = params.get("estado", [""])[0].strip()
        values: list[object] = []
        filters = []
        if departamento:
            filters.append("departamento = ?")
            values.append(departamento)
        if marca:
            filters.append("marca = ?")
            values.append(marca)
        if proveedor:
            filters.append("proveedor = ?")
            values.append(proveedor)
        stock_expr = "(COALESCE(stock_unidad, stock, 0) + COALESCE(stock_deposito, 0))"
        if estado == "sin_stock":
            filters.append(f"{stock_expr} <= 0")
        elif estado == "stock_bajo":
            filters.append(f"stock_minimo > 0 AND {stock_expr} <= stock_minimo")
        where = " WHERE " + " AND ".join(filters) if filters else ""

        conn = connect()
        resumen = execute(
            conn,
            f"""
            SELECT COUNT(*) AS productos,
                   COALESCE(SUM(COALESCE(stock_unidad, stock, 0)), 0) AS stock_unidad,
                   COALESCE(SUM(COALESCE(stock_deposito, 0)), 0) AS stock_deposito,
                   COALESCE(SUM({stock_expr}), 0) AS stock_total,
                   SUM(CASE WHEN {stock_expr} <= 0 THEN 1 ELSE 0 END) AS sin_stock,
                   SUM(CASE WHEN stock_minimo > 0 AND {stock_expr} <= stock_minimo THEN 1 ELSE 0 END) AS stock_bajo
            FROM productos
            {where}
            """,
            values,
        ).fetchone()
        por_departamento = [
            dict(row)
            for row in execute(
                conn,
                f"""
                SELECT COALESCE(NULLIF(TRIM(departamento), ''), 'Sin categoría') AS departamento,
                       COUNT(*) AS productos,
                       COALESCE(SUM({stock_expr}), 0) AS stock_total
                FROM productos
                {where}
                GROUP BY departamento
                ORDER BY stock_total DESC, productos DESC
                LIMIT 12
                """,
                values,
            ).fetchall()
        ]
        criticos = [
            normalized_stock(row)
            for row in execute(
                conn,
                f"""
                SELECT id, codigo_item, producto, marca, proveedor, departamento, descripcion,
                       COALESCE(stock_unidad, stock, 0) AS stock_unidad,
                       COALESCE(stock_deposito, 0) AS stock_deposito,
                       stock_minimo
                FROM productos
                WHERE {stock_expr} <= 0 OR (stock_minimo > 0 AND {stock_expr} <= stock_minimo)
                ORDER BY {stock_expr} ASC, codigo_item
                LIMIT 80
                """,
            ).fetchall()
        ]
        conn.close()
        json_response(self, {"resumen": dict(resumen), "por_departamento": por_departamento, "criticos": criticos})

    def parse_excel_rows(self, content: bytes) -> tuple[list[dict], list[str]]:
        if load_workbook is None:
            return [], ["openpyxl no está instalado"]
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        rows: list[dict] = []
        errors: list[str] = []
        for sheet in workbook.worksheets:
            raw_rows = sheet.iter_rows(values_only=True)
            headers = next(raw_rows, None)
            if not headers:
                continue
            names = [str(value or "").strip() for value in headers]
            for index, values in enumerate(raw_rows, start=2):
                row = {names[i]: values[i] if i < len(values) else "" for i in range(len(names))}
                if not any(str(value or "").strip() for value in row.values()):
                    continue
                codigo_item = str(pick(row, "Código Item", "Codigo Item", "codigo_item", "Código", "Codigo")).strip()
                codigo_barras = str(pick(row, "Código de Barras", "Codigo de Barras", "codigo_barras", "Barra")).strip()
                codigo_articulo = str(pick(row, "CodArticulo", "Cod Articulo", "Código Artículo", "codigo_articulo")).strip()
                product_name = str(pick(row, "Producto", "Descripción corta", "Articulo", "Artículo")).strip()
                if not (codigo_item or codigo_barras or codigo_articulo or product_name):
                    errors.append(f"{sheet.title} fila {index}: fila sin código ni producto")
                    continue
                rows.append(
                    {
                        "hoja": sheet.title,
                        "fila": index,
                        "codigo_item": codigo_item,
                        "codigo_barras": codigo_barras,
                        "codigo_articulo": codigo_articulo,
                        "producto": product_name,
                        "marca": str(pick(row, "Marca")).strip(),
                        "proveedor": str(pick(row, "Proveedor")).strip(),
                        "departamento": str(pick(row, "Departamento", "Categoría", "Categoria", "Rubro")).strip(),
                        "descripcion": str(pick(row, "Descripción", "Descripcion", "Detalle")).strip(),
                        "aplicacion": str(pick(row, "Aplicación", "Aplicacion", "Vehículo", "Vehiculo")).strip(),
                        "stock_unidad": to_int(pick(row, "Unidad", "Mostrador", "Stock Unidad", "Stock Mostrador")),
                        "stock_deposito": to_int(pick(row, "Depósito", "Deposito", "Stock Depósito", "Stock Deposito", "Stock")),
                    }
                )
        return rows, errors

    def handle_excel_import(self) -> None:
        user = self.require_user(role="administrador")
        if user is None:
            return
        data = self.read_json()
        filename = str(data.get("filename", "importacion.xlsx")).strip() or "importacion.xlsx"
        preview_only = bool(data.get("preview", True))
        content_b64 = str(data.get("content", ""))
        try:
            content = base64.b64decode(content_b64)
        except Exception:
            json_response(self, {"error": "Archivo inválido"}, 400)
            return
        rows, errors = self.parse_excel_rows(content)
        preview = rows[:30]
        if preview_only:
            json_response(self, {"preview": preview, "total_filas": len(rows), "errores": errors[:50]})
            return

        conn = connect()
        nuevos = actualizados = ignorados = 0
        try:
            begin_write(conn)
            for row in rows:
                stock_total = row["stock_unidad"] + row["stock_deposito"]
                lookup_values = [row["codigo_item"], row["codigo_barras"], row["codigo_articulo"]]
                existing = None
                for column, value in zip(["codigo_item", "codigo_barras", "codigo_articulo"], lookup_values):
                    if value:
                        existing = execute(conn, f"SELECT id, COALESCE(stock_unidad, stock, 0) AS stock_unidad, COALESCE(stock_deposito, 0) AS stock_deposito FROM productos WHERE {column} = ? LIMIT 1", (value,)).fetchone()
                        if existing:
                            break
                if existing:
                    previous = normalized_stock(existing)["stock_total"]
                    execute(
                        conn,
                        """
                        UPDATE productos
                        SET codigo_item = ?, codigo_barras = ?, codigo_articulo = ?, producto = ?,
                            marca = ?, proveedor = ?, departamento = ?, descripcion = ?, aplicacion = ?,
                            stock_unidad = ?, stock_deposito = ?, stock = ?, fecha_actualizacion = ?
                        WHERE id = ?
                        """,
                        (
                            row["codigo_item"], row["codigo_barras"], row["codigo_articulo"], row["producto"],
                            row["marca"], row["proveedor"], row["departamento"], row["descripcion"], row["aplicacion"],
                            row["stock_unidad"], row["stock_deposito"], stock_total, now_text(), existing["id"],
                        ),
                    )
                    self.register_movement(conn, existing["id"], "importacion_excel", stock_total - previous, previous, stock_total, f"Importación Excel: {filename}", user)
                    actualizados += 1
                else:
                    insert_sql = """
                        INSERT INTO productos (
                            archivo_origen, hoja_origen, fila_origen, codigo_item, codigo_barras,
                            codigo_articulo, producto, marca, proveedor, departamento, descripcion, aplicacion,
                            stock, stock_unidad, stock_deposito, fecha_actualizacion
                        )
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """
                    if IS_POSTGRES:
                        insert_sql += " RETURNING id"
                    cursor = execute(
                        conn,
                        insert_sql,
                        (
                            filename, row["hoja"], row["fila"], row["codigo_item"], row["codigo_barras"],
                            row["codigo_articulo"], row["producto"], row["marca"], row["proveedor"], row["departamento"],
                            row["descripcion"], row["aplicacion"], stock_total, row["stock_unidad"], row["stock_deposito"], now_text(),
                        ),
                    )
                    product_id = cursor.fetchone()["id"] if IS_POSTGRES else cursor.lastrowid
                    self.register_movement(conn, product_id, "importacion_excel", stock_total, 0, stock_total, f"Importación Excel: {filename}", user)
                    nuevos += 1
            conn.commit()
        except Exception as exc:
            conn.rollback()
            json_response(self, {"error": f"No se pudo importar: {exc}"}, 400)
            return
        finally:
            conn.close()
        json_response(
            self,
            {"ok": True, "nuevos": nuevos, "actualizados": actualizados, "ignorados": ignorados, "errores": errors[:50]},
        )


def main() -> None:
    init_db()
    host = os.environ.get("HOST", "0.0.0.0")
    port = int(os.environ.get("PORT", "8000"))
    server = ThreadingHTTPServer((host, port), AppHandler)
    print(f"Servidor iniciado en http://127.0.0.1:{port}")
    print(f"Escuchando en {host}:{port}")
    server.serve_forever()


if __name__ == "__main__":
    main()
