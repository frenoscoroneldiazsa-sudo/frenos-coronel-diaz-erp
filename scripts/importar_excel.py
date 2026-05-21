from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "EXCEL_ACTUALES"
OUTPUT_DIR = ROOT / "DATOS_IMPORTADOS"


OUTPUT_COLUMNS = [
    "archivo_origen",
    "hoja_origen",
    "fila_origen",
    "codigo_item",
    "codigo_barras",
    "codigo_articulo",
    "producto",
    "marca",
    "proveedor",
    "departamento",
    "descripcion",
    "unidad",
    "stock",
    "deposito",
    "entrada",
    "salida",
    "stock_calculado",
]


HEADER_ALIASES = {
    "codigo_item": {
        "codigo item",
        "codigoitem",
        "codigo iterm",
    },
    "codigo_barras": {
        "codigo de barras",
        "codigo barras",
        "codigo barra",
        "codigodebarras",
    },
    "codigo_articulo": {
        "codarticulo",
        "cod articulo",
        "codartituclo",
        "podarticulo",
    },
    "producto": {
        "producto",
        "prodcuto",
    },
    "marca": {
        "marca",
    },
    "proveedor": {
        "proveedor",
    },
    "departamento": {
        "departamento",
    },
    "descripcion": {
        "descripcion",
        "descripccion",
        "descripcion - medida",
    },
    "unidad": {
        "unidad",
    },
    "stock": {
        "stock",
    },
    "deposito": {
        "deposito",
    },
    "entrada": {
        "entrada",
    },
    "salida": {
        "salida",
    },
}


def clean_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.strip().lower()
    text = text.replace("_", " ")
    text = re.sub(r"\s+", " ", text)
    text = text.replace("codigo de barra", "codigo barra")
    return text


def clean_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_number(value: str) -> float:
    text = clean_value(value)
    if not text:
        return 0.0
    text = text.replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return 0.0
    return float(match.group(0))


def find_header_map(ws: Any) -> tuple[int | None, dict[int, str]]:
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True),
        start=1,
    ):
        raw_headers = [clean_header(value) for value in row]
        header_map: dict[int, str] = {}
        for col_idx, header in enumerate(raw_headers, start=1):
            if not header:
                continue
            for canonical, aliases in HEADER_ALIASES.items():
                if header in aliases:
                    header_map[col_idx] = canonical
                    break

        useful = set(header_map.values())
        if "codigo_item" not in useful and "codigo_barras" in useful and raw_headers and not raw_headers[0]:
            header_map[1] = "codigo_item"
            useful.add("codigo_item")

        if "codigo_item" in useful and ("producto" in useful or "descripcion" in useful):
            return row_idx, header_map

    return None, {}


def normalize_row(
    file_name: str,
    sheet_name: str,
    row_idx: int,
    row_values: list[Any],
    header_map: dict[int, str],
) -> dict[str, str]:
    item = {column: "" for column in OUTPUT_COLUMNS}
    item["archivo_origen"] = file_name
    item["hoja_origen"] = sheet_name
    item["fila_origen"] = str(row_idx)

    for col_idx, value in enumerate(row_values, start=1):
        field = header_map.get(col_idx)
        if field:
            item[field] = clean_value(value)

    if item["proveedor"].strip().upper() == "RT":
        item["proveedor"] = "CARXON"
        item["producto"] = "RT"

    base_stock = to_number(item["unidad"] or item["stock"])
    deposito = to_number(item["deposito"])
    entrada = to_number(item["entrada"])
    salida = to_number(item["salida"])

    # Many sheets already use UNIDAD as current stock. We keep original columns
    # and add a calculated fallback for later validation.
    item["stock_calculado"] = str(int(base_stock + deposito + entrada - salida))

    return item


def row_has_product_data(item: dict[str, str]) -> bool:
    return bool(
        item["codigo_item"]
        or item["codigo_barras"]
        or item["codigo_articulo"]
        or item["producto"]
        or item["descripcion"]
    )


def import_workbooks() -> dict[str, Any]:
    OUTPUT_DIR.mkdir(exist_ok=True)

    rows: list[dict[str, str]] = []
    sheet_summaries: list[dict[str, Any]] = []
    skipped_sheets: list[dict[str, str]] = []

    for workbook_path in sorted(INPUT_DIR.glob("*.xlsx")):
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        for ws in workbook.worksheets:
            header_row, header_map = find_header_map(ws)
            if header_row is None:
                skipped_sheets.append(
                    {"archivo": workbook_path.name, "hoja": ws.title, "motivo": "sin encabezado reconocible"}
                )
                continue

            sheet_count = 0
            for row_idx, row_values in enumerate(
                ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, values_only=True),
                start=header_row + 1,
            ):
                item = normalize_row(workbook_path.name, ws.title, row_idx, row_values, header_map)
                if row_has_product_data(item):
                    rows.append(item)
                    sheet_count += 1

            sheet_summaries.append(
                {
                    "archivo": workbook_path.name,
                    "hoja": ws.title,
                    "filas_importadas": sheet_count,
                    "encabezado_fila": header_row,
                    "columnas_detectadas": sorted(set(header_map.values())),
                }
            )

    csv_path = OUTPUT_DIR / "inventario_unificado.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    by_file = Counter(row["archivo_origen"] for row in rows)
    by_sheet = defaultdict(int)
    for row in rows:
        by_sheet[f'{row["archivo_origen"]} / {row["hoja_origen"]}'] += 1

    summary = {
        "generado": datetime.now().isoformat(timespec="seconds"),
        "archivos_leidos": len(list(INPUT_DIR.glob("*.xlsx"))),
        "hojas_importadas": len(sheet_summaries),
        "hojas_omitidas": skipped_sheets,
        "filas_importadas": len(rows),
        "productos_con_codigo_item": sum(1 for row in rows if row["codigo_item"]),
        "productos_con_codigo_barras": sum(1 for row in rows if row["codigo_barras"]),
        "productos_con_descripcion": sum(1 for row in rows if row["descripcion"]),
        "filas_por_archivo": dict(by_file),
        "filas_por_hoja": dict(sorted(by_sheet.items())),
        "detalle_hojas": sheet_summaries,
        "salida_csv": str(csv_path),
    }

    summary_path = OUTPUT_DIR / "resumen_importacion.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


if __name__ == "__main__":
    result = import_workbooks()
    print(json.dumps(result, ensure_ascii=False, indent=2))
